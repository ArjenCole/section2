"""汇总数据组装与 Excel 导出（计划 M5 / §7.3）。

* ``collect_summary``：对一组单位工程计算工程量，组装定额汇总与清单逐构件两套行，
  并按单价字典算合价（等价旧版 FormSum 的 Summary + FlashdGVQ/FlashdGVL）；
* ``export_excel``：写出 定额工程量 / 清单工程量 两个 sheet，
  列结构与旧版一致，并按 §7.3 在“工程量”列前插入“计算表达式”列。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.evaluator import format_number
from app.orm.models import CATEGORY_NAMES, Element
from app.services import project_io
from app.services.calc.tracer import QDict, sorted_items

#: 清单项目行的底色（旧版 LightBlue）
_LISTING_FILL = PatternFill("solid", fgColor="D6EAF8")
_HEADER_FILL = PatternFill("solid", fgColor="EFF6FF")


@dataclass
class QuantityRow:
    """定额 / 清单里的一行工程量。"""

    no: str = ""  # 编号（清单项目行放“清单项目”字样）
    category: str = ""
    item: str = ""
    unit: str = ""
    expression: str = ""
    amount: float = 0.0
    price: float = 0.0
    total: float = 0.0
    editable_price: bool = False


@dataclass
class SummaryData:
    """汇总结果：定额行 + 清单行（清单项目行与其工程量行交错）+ 分类合计。"""

    quota: list[QuantityRow] = field(default_factory=list)
    listing: list[QuantityRow] = field(default_factory=list)
    quota_category_totals: dict[str, float] = field(default_factory=dict)
    listing_category_totals: dict[str, float] = field(default_factory=dict)
    quota_total: float = 0.0
    listing_total: float = 0.0


def _key_parts(key: str) -> tuple[str, str, str]:
    parts = key.split("|")
    return (
        parts[0] if parts else "",
        parts[1] if len(parts) > 1 else "",
        parts[2] if len(parts) > 2 else "",
    )


def _quantity_row(key: str, quantity, with_price: bool) -> QuantityRow:
    category, item, unit = _key_parts(key)
    price = project_io.get_price(key) if with_price else 0.0
    return QuantityRow(
        category=category,
        item=item,
        unit=unit,
        expression=quantity.expression,
        amount=quantity.value,
        price=price,
        total=price * quantity.value,
        editable_price=with_price,
    )


def collect_summary(unit_ids: list[int] | None = None) -> SummaryData:
    """计算并组装汇总（unit_ids 为空时取全部单位工程）。"""
    from app.viewmodels.unit_vm import compute_unit

    data = SummaryData()
    quota = QDict()
    if unit_ids is None:
        unit_ids = [
            unit.id
            for segment in project_io.segments()
            for unit in project_io.units(segment.id)
        ]
    for unit_id in unit_ids:
        _summary, detail = compute_unit(unit_id)
        for row, element_dq in detail:
            data.listing.append(
                QuantityRow(
                    no="清单项目",
                    category=CATEGORY_NAMES.get(row.category, ""),
                    item=row.spec or row.name,
                    unit=row.unit,
                    amount=row.amount_value,
                )
            )
            subtotal = 0.0
            for quantity in sorted_items(element_dq):
                listing_row = _quantity_row(quantity.key, quantity, with_price=True)
                data.listing.append(listing_row)
                subtotal += listing_row.total
                category_total = data.listing_category_totals.get(listing_row.category, 0.0)
                data.listing_category_totals[listing_row.category] = category_total + listing_row.total
            data.listing_total += subtotal
        quota.extend(_summary)

    for quantity in sorted_items(quota):
        row = _quantity_row(quantity.key, quantity, with_price=True)
        data.quota.append(row)
        data.quota_total += row.total
        data.quota_category_totals[row.category] = (
            data.quota_category_totals.get(row.category, 0.0) + row.total
        )
    return data


def _footer_rows(data: SummaryData, table: str) -> list[QuantityRow]:
    totals = data.quota_category_totals if table == "quota" else data.listing_category_totals
    grand = data.quota_total if table == "quota" else data.listing_total
    rows = [QuantityRow(no="合计：", unit="元", total=grand)]
    for category, value in totals.items():
        rows.append(QuantityRow(category=category, unit="元", total=value))
    return rows


def export_excel(path: str | Path, data: SummaryData) -> bool:
    """写出双表 Excel（定额工程量 / 清单工程量）。"""
    workbook = Workbook()
    quota_sheet = workbook.active
    quota_sheet.title = "定额工程量"
    _write_sheet(quota_sheet, data.quota, _footer_rows(data, "quota"))
    listing_sheet = workbook.create_sheet("清单工程量")
    _write_sheet(listing_sheet, data.listing, _footer_rows(data, "listing"))
    workbook.save(str(path))
    return True


_HEADERS = ["编号", "类别", "项目", "单位", "计算表达式", "工程量", "单价", "合价"]


def _write_sheet(sheet, rows: list[QuantityRow], footers: list[QuantityRow]) -> None:
    bold = Font(bold=True)
    for column, title in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = bold
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    row_index = 2
    for row in rows:
        if row.no == "清单项目":
            for column in range(1, len(_HEADERS) + 1):
                sheet.cell(row=row_index, column=column).fill = _LISTING_FILL
        values = (
            row.no,
            row.category,
            row.item,
            row.unit,
            row.expression,
            format_number(row.amount) if row.amount else "",
            format_number(row.price) if row.price else "",
            format_number(row.total) if row.total else "",
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)
        row_index += 1
    for footer in footers:
        sheet.cell(row=row_index, column=1, value=footer.no)
        sheet.cell(row=row_index, column=2, value=footer.category)
        sheet.cell(row=row_index, column=4, value=footer.unit)
        cell = sheet.cell(row=row_index, column=8, value=format_number(footer.total))
        cell.font = bold
        row_index += 1
    widths = (8, 14, 30, 8, 46, 12, 10, 12)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
